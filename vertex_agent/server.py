from __future__ import annotations

import asyncio
import json
import os
import shutil
import socket
import subprocess
import sys
import time
import uuid
from contextlib import asynccontextmanager
from io import BytesIO
from pathlib import Path

import httpx
from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from vertex_agent import __version__
from vertex_agent.agent import run_turn
from vertex_agent.config import (
    assert_workspace_allowed,
    is_page_model,
    load_config,
    model_base,
    model_for_task,
    save_runtime_settings,
)
from vertex_agent.deepseek import verify_api_key
from vertex_agent.llm import ChatRouter, ModelError
from vertex_agent.secrets import (
    MODEL_ID,
    clear_api_key,
    is_deepseek_model,
    is_loopback,
    key_status,
    save_api_key,
)
from vertex_agent.memory import MemoryStore, SessionStore
from vertex_agent.prompts import CURSOR_PROMPT
from vertex_agent.tools import Toolbox

STATIC_DIR = Path(__file__).resolve().parent / "static"
locks: dict[str, asyncio.Lock] = {}


class ChatIn(BaseModel):
    message: str
    session_id: str | None = None
    mode: str = "general"
    think: bool = False
    images: list[str] = Field(default_factory=list)
    language: str = ""
    model: str | None = None


class SettingsIn(BaseModel):
    model: str | None = None
    general_model: str | None = None
    workspace: str | None = None


class FileBody(BaseModel):
    path: str
    content: str = ""


class RunBody(BaseModel):
    path: str
    action: str = "run"
    args: list[str] = Field(default_factory=list)


class SourceBody(BaseModel):
    path: str = ""
    content: str = ""
    scope: str = "agent"
    action: str = "compile"
    args: list[str] = Field(default_factory=list)
    note: str = ""


class KeyIn(BaseModel):
    api_key: str


class CompletionIn(BaseModel):
    model: str | None = None
    messages: list[dict] = Field(default_factory=list)
    stream: bool = False
    tools: list[dict] | None = None
    temperature: float | None = None


@asynccontextmanager
async def lifespan(app: FastAPI):
    config = load_config()
    config.workspace.mkdir(parents=True, exist_ok=True)
    (config.workspace / "_runs").mkdir(parents=True, exist_ok=True)
    config.data_dir.mkdir(parents=True, exist_ok=True)
    memory = MemoryStore(config.data_dir / "memory.json")
    app.state.config = config
    app.state.memory = memory
    app.state.sessions = SessionStore(config.data_dir / "sessions")
    app.state.ollama = ChatRouter(config.base_url, config.timeout_sec, config.data_dir)
    app.state.toolbox = Toolbox(
        config.project_root,
        config.workspace,
        config.data_dir,
        memory,
        matlab=config.matlab,
        self_edit=config.self_edit,
    )
    yield
    await app.state.ollama.aclose()


app = FastAPI(title="VertexAgent", version=__version__, lifespan=lifespan)
class NoCacheStatic(StaticFiles):
    async def get_response(self, path: str, scope):
        response = await super().get_response(path, scope)
        response.headers["Cache-Control"] = "no-cache"
        return response


app.mount("/static", NoCacheStatic(directory=STATIC_DIR), name="static")


@app.get("/")
async def index():
    return FileResponse(STATIC_DIR / "index.html", headers={"Cache-Control": "no-cache"})


@app.get("/api/health")
async def health():
    config = app.state.config
    ollama_ok = False
    try:
        await app.state.ollama.tags()
        ollama_ok = True
    except (httpx.HTTPError, OSError):
        ollama_ok = False
    return {
        "name": "VertexAgent",
        "version": __version__,
        "model": config.model,
        "general_model": config.general_model,
        "ollama": ollama_ok,
        "self_edit": config.self_edit,
        "workspace": str(config.workspace),
        "local": f"http://127.0.0.1:{config.port}",
        "lan": lan_urls(config.port),
        "gcc": bool(shutil.which("gcc")),
        "matlab": bool(config.matlab and Path(config.matlab).exists()),
        "deepseek": key_status(config.data_dir),
        "python": f"{sys.version_info.major}.{sys.version_info.minor}.{sys.version_info.micro}",
    }


@app.get("/api/deepseek")
async def deepseek_status():
    return key_status(app.state.config.data_dir)


@app.post("/api/deepseek")
async def deepseek_register(body: KeyIn, request: Request):
    if not is_loopback(request.client.host if request.client else ""):
        raise HTTPException(status_code=403, detail="请在这台电脑打开 http://127.0.0.1:8765 后再注册 Key")
    api_key = body.api_key.strip()
    if not api_key.startswith("sk-") or len(api_key) < 20 or any(ch.isspace() for ch in api_key):
        raise HTTPException(status_code=400, detail="Key 应以 sk- 开头，且不能包含空格")
    try:
        await verify_api_key(api_key)
    except ModelError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    save_api_key(app.state.config.data_dir, api_key)
    status = key_status(app.state.config.data_dir)
    return {"ok": True, "model": app.state.config.model, **status}


@app.delete("/api/deepseek")
async def deepseek_clear(request: Request):
    if not is_loopback(request.client.host if request.client else ""):
        raise HTTPException(status_code=403, detail="只能在这台电脑上清除 Key")
    clear_api_key(app.state.config.data_dir)
    if is_deepseek_model(app.state.config.model):
        app.state.config.model = "qwen3.5:4b"
        app.state.config.general_model = "qwen3.5:4b"
        save_runtime_settings({"model": "qwen3.5:4b", "general_model": "qwen3.5:4b"})
    return {"ok": True, "registered": False, "model": app.state.config.model}


@app.get("/api/models")
async def list_models():
    try:
        listed = await app.state.ollama.tags()
    except (httpx.HTTPError, OSError):
        listed = []
    qwen_name = "qwen3.5:4b"
    for item in listed:
        name = item.get("name") or item.get("model") or ""
        if model_base(name).startswith("qwen3.5"):
            qwen_name = name
            break
    models = [
        {"name": "deepseek-v4-pro", "label": "DeepSeek V4.1 Pro"},
        {"name": MODEL_ID, "label": "DeepSeek V4.1 Flash"},
        {"name": qwen_name, "label": "qwen3.5"},
    ]
    active = app.state.config.model
    if not is_page_model(active):
        active = qwen_name
        app.state.config.model = active
        app.state.config.general_model = active
        save_runtime_settings({"model": active, "general_model": active})
    return {"models": models, "active": active, "general_model": app.state.config.general_model}


@app.post("/api/settings")
async def update_settings(body: SettingsIn):
    config = app.state.config
    updates = {}
    if body.model is not None:
        model = body.model.strip()
        if not model or len(model) > 80 or any(ch in model for ch in "\\/\n\r"):
            raise HTTPException(status_code=400, detail="模型名称不合法")
        if not is_page_model(model):
            raise HTTPException(status_code=400, detail="模型只支持 DeepSeek V4.1 Pro、DeepSeek V4.1 Flash 和 qwen3.5")
        config.model = model
        updates["model"] = model
        config.general_model = model
        updates["general_model"] = model
    if body.general_model is not None:
        general = body.general_model.strip()
        if general and is_page_model(general):
            config.general_model = general
            updates["general_model"] = general
    if body.workspace is not None:
        try:
            workspace = assert_workspace_allowed(Path(body.workspace), config.project_root)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        workspace.mkdir(parents=True, exist_ok=True)
        (workspace / "_runs").mkdir(parents=True, exist_ok=True)
        config.workspace = workspace
        app.state.toolbox.workspace = workspace
        updates["workspace"] = str(workspace)
    if updates:
        save_runtime_settings(updates)
    return {"ok": True, "model": config.model, "workspace": str(config.workspace)}


@app.get("/api/sessions")
async def list_sessions():
    return {"sessions": app.state.sessions.list_sessions()}


@app.get("/api/sessions/{session_id}")
async def get_session(session_id: str):
    try:
        data = app.state.sessions.load(session_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if data is None:
        raise HTTPException(status_code=404, detail="会话不存在")
    return data


@app.delete("/api/sessions/{session_id}")
async def delete_session(session_id: str):
    try:
        deleted = app.state.sessions.delete(session_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not deleted:
        raise HTTPException(status_code=404, detail="会话不存在")
    return {"ok": True}


@app.post("/api/chat")
async def chat(body: ChatIn):
    text = body.message.strip()
    images = _clean_images(body.images)
    if not text and not images:
        raise HTTPException(status_code=400, detail="消息为空")
    if not text:
        text = "请看这张图片。"
    if body.mode not in {"general", "workspace", "self"}:
        raise HTTPException(status_code=400, detail="未知模式")
    try:
        model = model_for_task(
            app.state.config.model,
            body.model,
            language=body.language,
            mode=body.mode,
            general_model=app.state.config.general_model,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    session_id = body.session_id or f"s-{uuid.uuid4().hex[:12]}"
    lock = locks.setdefault(session_id, asyncio.Lock())

    async def generate():
        async with lock:
            async for event in run_turn(
                ollama=app.state.ollama,
                config=app.state.config,
                toolbox=app.state.toolbox,
                sessions=app.state.sessions,
                memory=app.state.memory,
                session_id=session_id,
                text=text,
                mode=body.mode,
                think=body.think,
                images=images,
                model=model,
                language=body.language,
            ):
                yield _sse(event)

    return StreamingResponse(generate(), media_type="text/event-stream", headers={"Cache-Control": "no-cache"})


def _extract_document(name: str, data: bytes) -> str:
    if len(data) > 16 * 1024 * 1024:
        raise ValueError("单个文件不要超过 16MB。旧版 doc、ppt 请先另存为 docx、pptx。")
    suffix = Path(name or "").suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    elif suffix == ".docx":
        from docx import Document

        text = "\n".join(paragraph.text for paragraph in Document(BytesIO(data)).paragraphs)
    elif suffix == ".pptx":
        from pptx import Presentation

        lines = []
        for slide in Presentation(BytesIO(data)).slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    lines.append(shape.text_frame.text)
        text = "\n".join(lines)
    else:
        text = data.decode("utf-8", errors="replace")
    text = (text or "").strip()
    if not text:
        raise ValueError("没有读到正文")
    return text[:20000]


@app.post("/api/import")
async def import_document(file: UploadFile = File(...)):
    data = await file.read()
    try:
        text = _extract_document(file.filename or "", data)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail="这个文件读不出来，请把正文复制到输入框") from exc
    return {"ok": True, "name": file.filename or "", "text": text}


@app.get("/api/files")
async def files(scope: str = "workspace", path: str = ""):
    return json.loads(app.state.toolbox.execute("list_files", {"scope": scope, "path": path}))


@app.get("/api/workspace/raw")
async def workspace_raw(path: str = ""):
    if not path:
        raise HTTPException(status_code=400, detail="缺少路径")
    try:
        target = app.state.toolbox._file("workspace", path, must_exist=True)
        from vertex_agent.safety import read_text

        return {"ok": True, "path": path, "content": read_text(target)}
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/workspace/save")
async def workspace_save(body: FileBody):
    try:
        return json.loads(app.state.toolbox.save_text(body.path, body.content))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/workspace/run")
async def workspace_run(body: RunBody):
    if body.action not in {"compile", "debug", "run", "verify"}:
        raise HTTPException(status_code=400, detail="动作必须是 compile、debug、run 或 verify")
    try:
        if body.action == "verify":
            return json.loads(app.state.toolbox.build_action("workspace", body.path, "verify", body.args))
        return json.loads(app.state.toolbox.run_saved(body.path, body.action, body.args))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/api/source/raw")
async def source_raw(scope: str = "agent", path: str = ""):
    if scope not in {"workspace", "agent"} or not path:
        raise HTTPException(status_code=400, detail="缺少路径")
    try:
        return json.loads(app.state.toolbox.read_raw(scope, path))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/source/save")
async def source_save(body: SourceBody):
    try:
        return json.loads(app.state.toolbox.save_scoped(body.scope, body.path, body.content))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/source/check")
async def source_check(body: SourceBody):
    try:
        return json.loads(app.state.toolbox.build_action(body.scope, body.path, "compile", body.args))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/source/action")
async def source_action(body: SourceBody):
    if body.action not in {"compile", "debug", "run", "verify"}:
        raise HTTPException(status_code=400, detail="动作必须是 compile、debug、run 或 verify")
    try:
        return json.loads(app.state.toolbox.build_action(body.scope, body.path, body.action, body.args))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/source/publish")
async def source_publish(body: SourceBody):
    try:
        return json.loads(app.state.toolbox.publish_release(body.note))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/api/file")
async def file(scope: str = "workspace", path: str = ""):
    if not path:
        raise HTTPException(status_code=400, detail="缺少路径")
    return json.loads(app.state.toolbox.execute("read_file", {"scope": scope, "path": path, "end_line": 400}))


@app.get("/api/self-edits")
async def self_edits():
    return {"edits": app.state.toolbox.log.list_entries(50)}


@app.post("/api/self-edits/{edit_id}/rollback")
async def rollback(edit_id: str):
    try:
        return app.state.toolbox.rollback(edit_id)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/restart")
async def restart():
    project = str(app.state.config.project_root)
    flags = 0
    if os.name == "nt":
        flags = subprocess.CREATE_NEW_PROCESS_GROUP | subprocess.DETACHED_PROCESS | subprocess.CREATE_NO_WINDOW
    subprocess.Popen(
        [sys.executable, "-c", "import time,subprocess,sys; time.sleep(1.2); subprocess.Popen([sys.executable,'-m','vertex_agent'])"],
        cwd=project,
        creationflags=flags,
        close_fds=True,
    )

    async def _exit_soon():
        await asyncio.sleep(0.4)
        os._exit(0)

    asyncio.create_task(_exit_soon())
    return {"ok": True, "message": "服务即将重启"}


@app.get("/v1/models")
async def openai_models():
    return {
        "object": "list",
        "data": [{"id": "vertex-agent", "object": "model", "owned_by": "vertex", "created": int(time.time())}],
    }


@app.post("/v1/chat/completions")
async def openai_chat(body: CompletionIn):
    messages = _adapt_messages(body.messages)
    temperature = body.temperature if body.temperature is not None else app.state.config.temperature
    config = app.state.config

    async def collect():
        turn = None
        async for event in app.state.ollama.chat(
            model=config.model,
            messages=messages,
            tools=body.tools,
            temperature=temperature,
            num_ctx=config.num_ctx,
            num_predict=config.num_predict,
            think=False,
        ):
            if event["type"] == "turn":
                turn = event
            elif body.stream and event["type"] == "token":
                yield event
        if turn is None:
            raise ModelError("模型没有返回内容")
        if body.stream:
            yield {"type": "turn", **turn}
        else:
            yield turn

    if not body.stream:
        try:
            turn = None
            async for item in collect():
                turn = item
        except ModelError as exc:
            raise HTTPException(status_code=502, detail=str(exc)) from exc
        return _openai_message(turn or {})

    async def generate():
        completion_id = f"chatcmpl-{uuid.uuid4().hex[:12]}"
        created = int(time.time())
        try:
            async for event in collect():
                if event["type"] == "token" and event.get("text"):
                    yield _sse(_openai_chunk(completion_id, created, {"content": event["text"]}))
                elif event["type"] == "turn":
                    calls = _openai_tool_calls(event.get("tool_calls") or [])
                    if calls:
                        for index, call in enumerate(calls):
                            delta_call = dict(call)
                            delta_call["index"] = index
                            yield _sse(_openai_chunk(completion_id, created, {"tool_calls": [delta_call]}))
                        yield _sse(_openai_chunk(completion_id, created, {}, "tool_calls"))
                    else:
                        yield _sse(_openai_chunk(completion_id, created, {}, "stop"))
            yield "data: [DONE]\n\n"
        except ModelError as exc:
            yield _sse({"error": {"message": str(exc)}})

    return StreamingResponse(generate(), media_type="text/event-stream")


def _clean_images(images: list[str]) -> list[str]:
    cleaned = []
    for item in images[:4]:
        if isinstance(item, str) and item.startswith("data:image/") and len(item) <= 6_000_000:
            cleaned.append(item)
    return cleaned


def lan_urls(port: int) -> list[str]:
    urls = []
    try:
        infos = socket.getaddrinfo(socket.gethostname(), None, socket.AF_INET)
    except OSError:
        infos = []
    for info in infos:
        ip = info[4][0]
        if ip.startswith("127.") or ip.startswith("169.254."):
            continue
        urls.append(f"http://{ip}:{port}")
    return list(dict.fromkeys(urls))


def _sse(payload: dict) -> str:
    return f"data: {json.dumps(payload, ensure_ascii=False)}\n\n"


def _adapt_messages(messages: list[dict]) -> list[dict]:
    names: dict[str, str] = {}
    adapted = []
    for original in messages:
        message = dict(original)
        content = message.get("content")
        if isinstance(content, list):
            texts = []
            for part in content:
                if isinstance(part, str):
                    texts.append(part)
                elif isinstance(part, dict):
                    texts.append(str(part.get("text") or ""))
            message["content"] = "".join(texts)
        elif content is None:
            message["content"] = ""
        if message.get("role") == "assistant":
            calls = []
            for call in message.get("tool_calls") or []:
                item = dict(call)
                function = dict(item.get("function") or {})
                arguments = function.get("arguments")
                if isinstance(arguments, str):
                    try:
                        function["arguments"] = json.loads(arguments) if arguments.strip() else {}
                    except json.JSONDecodeError:
                        function["arguments"] = {"_raw": arguments}
                item["function"] = function
                if item.get("id") and function.get("name"):
                    names[item["id"]] = function["name"]
                calls.append(item)
            if calls:
                message["tool_calls"] = calls
        if message.get("role") == "tool" and not message.get("tool_name"):
            message["tool_name"] = names.get(message.get("tool_call_id") or "", "") or message.get("name") or ""
        adapted.append(message)
    if adapted and adapted[0].get("role") == "system":
        adapted[0]["content"] = CURSOR_PROMPT + "\n\n" + (adapted[0].get("content") or "")
    else:
        adapted.insert(0, {"role": "system", "content": CURSOR_PROMPT})
    return adapted


def _openai_tool_calls(calls: list[dict]) -> list[dict]:
    output = []
    for index, call in enumerate(calls):
        function = call.get("function") or {}
        arguments = function.get("arguments")
        if not isinstance(arguments, str):
            arguments = json.dumps(arguments or {}, ensure_ascii=False)
        output.append(
            {
                "id": call.get("id") or f"call_{index}",
                "type": "function",
                "function": {"name": function.get("name") or "", "arguments": arguments},
            }
        )
    return output


def _openai_message(turn: dict) -> dict:
    calls = _openai_tool_calls(turn.get("tool_calls") or [])
    message: dict = {"role": "assistant", "content": turn.get("content") or ""}
    if calls:
        message["tool_calls"] = calls
    return {
        "id": f"chatcmpl-{uuid.uuid4().hex[:12]}",
        "object": "chat.completion",
        "created": int(time.time()),
        "model": "vertex-agent",
        "choices": [
            {
                "index": 0,
                "message": message,
                "finish_reason": "tool_calls" if calls else "stop",
            }
        ],
    }


def _openai_chunk(completion_id: str, created: int, delta: dict, finish: str | None = None) -> dict:
    return {
        "id": completion_id,
        "object": "chat.completion.chunk",
        "created": created,
        "model": "vertex-agent",
        "choices": [{"index": 0, "delta": delta, "finish_reason": finish}],
    }
